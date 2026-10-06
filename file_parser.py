from pypdf import PdfReader
from docx import Document


def extract_text(file_path: str) -> str:
    """Read a PDF or DOCX and return plain text."""
    lower = file_path.lower()

    if lower.endswith(".pdf"):
        reader = PdfReader(file_path)
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages)

    elif lower.endswith(".docx"):
        doc = Document(file_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        # Many resumes use tables for layout, so read those too
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += "\n" + cell.text

    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    text = text.strip()
    if len(text) < 50:
        raise ValueError(
            "Could not read text from this file. It may be a scanned image."
        )
    return text